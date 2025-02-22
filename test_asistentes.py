import csv
import os
import unittest
from openpyxl import Workbook, load_workbook
from alumnosCsvToXslx_alfabeticamenteAbajo import escribir_excel

class TestAsistentes(unittest.TestCase):

    def setUp(self):
        self.ruta_csv = "test_asistentes.csv"
        self.ruta_excel ="test_asistentes.xlsx"
        self.datos_prueba =[
            ["Juan", "Pérez"],
            ["Ana", "Gómez"],
            ["Luis", "Martínez"],
            ["María", "López"]
        ]
        with open(self.ruta_csv, mode="w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            writer.writerows(self.datos_prueba)
        
    def tearDown(self):
        if os.path.exists(self.ruta_csv):
            os.remove(self.ruta_csv)
        if os.path.exists(self.ruta_excel):
            os.remove(self.ruta_excel)

    def test_escribir_excel(self):
        # tomar los datos, llamar a la función y escribir el excel
        workbook = Workbook()
        escribir_excel(self.ruta_excel, self.datos_prueba, 8)

        # comprobar que el excel existe
        self.assertTrue(os.path.exists(self.ruta_excel))

        # leer el excel si existe
        workbook = load_workbook(self.ruta_excel)

        # comprobar que los datos que hay en el excell se corresponden con los datos de prueba
        hoja = workbook.active
        self.assertEqual(hoja.cell(row=1, column=1).value, "Juan")
        self.assertEqual(hoja.cell(row=1, column=2).value, "Pérez")
    
    def test_escribir_excel_desde_datos(self):
        # tomar los datos, llamar a la función y escribir el excel
        workbook = Workbook()
        escribir_excel(self.ruta_excel, self.datos_prueba, 8)

        # comprobar que el excel existe
        self.assertTrue(os.path.exists(self.ruta_excel))

        # leer el excel si existe
        workbook = load_workbook(self.ruta_excel)

        # comprobar que los datos que hay en el excell se corresponden con los datos de prueba
        hoja = workbook.active
        alumno = self.datos_prueba[0]
        self.assertEqual(hoja.cell(row=1, column=1).value, alumno[0])
        self.assertEqual(hoja.cell(row=1, column=2).value, alumno[1])

    def test_escribir_excel_iterativo_segun_datos(self):
        # tomar los datos, llamar a la función y escribir el excel
        workbook = Workbook()
        escribir_excel(self.ruta_excel, self.datos_prueba, 8)

        # comprobar que el excel existe
        self.assertTrue(os.path.exists(self.ruta_excel))

        # leer el excel si existe
        workbook = load_workbook(self.ruta_excel)

        # comprobar que los datos que hay en el excell se corresponden con los datos de prueba
        hoja = workbook.active

        # Iterar por todos los datos
        alumnado = self.datos_prueba
        contador = 1
        for alumno in alumnado:
            self.assertEqual(hoja.cell(row=contador, column=1).value, alumno[0])# comprobar nombre
            self.assertEqual(hoja.cell(row=contador, column=2).value, alumno[1])# comprobar apellidos
            contador += 1


if __name__ == "__main__":
    unittest.main()