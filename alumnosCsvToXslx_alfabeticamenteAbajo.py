import csv
from openpyxl import Workbook

def mover_puntero(fila=0, columna=1, contador_personas=1, longitud=16, row_batch=8):
    fila += 1       

    if contador_personas % row_batch == 0:
        fila += 1

        if columna != 4 and (longitud - contador_personas*2) < row_batch:
            columna = 4
            fila = 1

    return {"fila":fila, "columna": columna, "contador_personas":contador_personas}

#función para escribir los excels desde una ruta
def escribir_excel(ruta:str, alumnado:list, row_batch = 8):
    # Creamos el archivo xlsx
    workbook = Workbook()

    # Seleccionamos la hoja activa
    hoja = workbook.active

    puntero = mover_puntero() #inicializar puntero

    # Iteramos sobre los nombres y apellidos
    for i in range(len(alumnado)):
        hoja.cell(row=puntero.get('fila'), column=puntero.get('columna')).value = alumnado[i][0] # Nombre
        hoja.cell(row=puntero.get('fila'), column=puntero.get('columna')+1).value = alumnado[i][1] # Apellido

        puntero = mover_puntero(puntero.get("fila"), puntero.get("columna"), i+1, 
                                len(alumnado), 8)        

    workbook.save(ruta)


# Leer el archivo csv y guardarlo en una lista
nombres_y_apellidos = []
with open("asistentes.csv", newline="", encoding="utf-8") as file:
    reader = csv.reader(file)
    for row in reader:
        nombres_y_apellidos.append(row)

# Ordenar los alumnos por apellidos
nombres_y_apellidos = sorted(nombres_y_apellidos, key=lambda x: (x[1], x[0]))

# Creamos el archivo xlsx
escribir_excel("asistentes.xlsx", nombres_y_apellidos, 8)